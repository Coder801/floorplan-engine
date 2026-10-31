"""Смета электрики: объёмы — из модели (код), цены — из прайса проекта (prices.yaml).

Чертежа нет — только таблица: out/estimate.md, out/estimate.csv и out/estimate.xlsx
(Excel с формулами: поменяли цену или количество — суммы пересчитываются).
Объёмы берутся из того же расчёта, что лист проводки (wiring.py); отдельной спецификации
изделий нет — смета без прайса и есть ведомость объёмов:

- кабель по маркам — с запасом RESERVE; работа по прокладке — по фактической длине;
- подрозетники и механизмы — по постам настенных механизмов, рамки — по блокам;
- распаечные коробки (во влажных помещениях — IP44, отдельная позиция); соединения под
  опрессовку — в каждой коробке L, N, PE и по одному на каждую клавишу выключателей,
  которые она питает;
- крепления кабеля — шаг FIXING_STEP по длине трасс;
- штробы — wall_chases (общие для спусков рядом): узкие (до CHASE_NARROW кабелей) и
  широкие — разные позиции; проходы через стены — wall_holes;
- щит: корпус и аппараты — по составу листа «Щит» (panel.build_panel: вид и номинал, как
  в таблице аппаратов на закупку), сборка — по модулям этого состава; щита нет — автомат
  на каждую группу; цены аппаратов — по виду (breaker, diff, rcd, input…), номинал —
  уточнением ключа (breaker.C16);
- светильники, выводы под подсветку и к технике — по точкам, LED-лента — в метрах;
- маркировка — бирка на каждом конце каждой линии; замер сопротивления изоляции с
  протоколом — на каждую силовую группу.

Цены не выдумываются: позиции без цены в прайсе остаются в смете с пометкой «нет цены»
и в итог не входят. Профиль и блоки питания LED в модели не описаны и в смету не входят.
"""

from __future__ import annotations

import csv
import math
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import Field

from floorplan_engine.electrical.boxes import BOX_ORDER, BOX_TITLES
from floorplan_engine.electrical.panel import BoardPanel, device_name
from floorplan_engine.electrical.wiring import RESERVE
from floorplan_engine.electrical.wiring_model import Wiring
from floorplan_engine.model.panel_catalog import DeviceKind
from floorplan_engine.model.schema import Electrical, PointType, StrictModel
from floorplan_engine.model.texts import Language, tr

PRICES_FILE = "prices.yaml"  # прайс проекта — в папке проекта, рядом с room.yaml
ESTIMATE_MD = "estimate.md"
ESTIMATE_CSV = "estimate.csv"
ESTIMATE_XLSX = "estimate.xlsx"
XLSX_MONEY = "#,##0.00"
XLSX_QTY = "#,##0.0##"
XLSX_EDIT_FILL = "FFF4D6"  # ячейки, которые правят руками: цена и количество
FIXING_STEP = 400.0  # мм: шаг креплений кабеля по трассе
SWITCHES = frozenset({PointType.SWITCH, PointType.PASS_SWITCH})
# Механизмы в подрозетнике: на каждый пост — подрозетник, на блок — рамка
MOUNTED = frozenset(
    {PointType.SOCKET, PointType.SWITCH, PointType.PASS_SWITCH, PointType.DATA, PointType.TV}
)

Section = Literal["material", "work"]

# Позиции сметы: ключ → раздел, единица, наименование (ru, uk). Ключи с параметром:
# cable.<марка>, frame.<постов>, junction_box.<размер>, junction_box_ip.<размер>, аппараты
# щита — <вид>.<номинал> (breaker.C16); нет цены с параметром — берётся цена без него
LINES: dict[str, tuple[Section, str, str, str]] = {
    "cable": ("material", "м", "Кабель {mark}", "Кабель {mark}"),
    "socket_box": ("material", "шт", "Подрозетник", "Підрозетник"),
    "junction_box": ("material", "шт", "Коробка распаечная {mark}", "Коробка розподільча {mark}"),
    "junction_box_ip": (
        "material",
        "шт",
        "Коробка распаечная IP44 {mark}",
        "Коробка розподільча IP44 {mark}",
    ),
    "crimp": (
        "material",
        "шт",
        "Гильза и термоусадка (опрессовка)",
        "Гільза і термоусадка (опресування)",
    ),
    "cable_fixing": ("material", "шт", "Крепление кабеля", "Кріплення кабелю"),
    "board_enclosure": ("material", "шт", "Корпус щита", "Корпус щита"),
    # Аппараты щита — по видам (DeviceKind); в строке — название и номинал из состава щита
    "breaker": ("material", "шт", "Автомат 1P (на группу)", "Автомат 1P (на групу)"),
    "diff": ("material", "шт", "Дифавтомат 1P+N", "Дифавтомат 1P+N"),
    "rcd": ("material", "шт", "УЗО", "ПЗВ"),
    "input": ("material", "шт", "Вводной автомат", "Вступний автомат"),
    "relay": ("material", "шт", "Реле напряжения", "Реле напруги"),
    "spd": ("material", "шт", "УЗИП", "ПЗІП"),
    "changeover": ("material", "шт", "Переключатель резерва", "Перемикач резерву"),
    "ats": ("material", "шт", "АВР", "АВР"),
    "feed": ("material", "шт", "Автомат питания инвертора", "Автомат живлення інвертора"),
    "socket": ("material", "шт", "Розетка с заземлением", "Розетка із заземленням"),
    "switch1": ("material", "шт", "Выключатель одноклавишный", "Вимикач одноклавішний"),
    "switch2": ("material", "шт", "Выключатель двухклавишный", "Вимикач двоклавішний"),
    "pass_switch": ("material", "шт", "Выключатель проходной", "Вимикач прохідний"),
    "data": ("material", "шт", "Розетка интернет", "Розетка інтернет"),
    "tv": ("material", "шт", "Розетка TV", "Розетка TV"),
    "frame": ("material", "шт", "Рамка {n}-постовая", "Рамка {n}-постова"),
    # Светильники и выводы — названия как в условных обозначениях (texts.symbol_*)
    "light": ("material", "шт", "", ""),
    "downlight": ("material", "шт", "", ""),
    "pendant": ("material", "шт", "", ""),
    "sconce": ("material", "шт", "", ""),
    "led_output": ("material", "шт", "", ""),
    "outlet": ("material", "шт", "", ""),
    "led_strip": ("material", "м", "", ""),
    "work.cable": ("work", "м", "Прокладка кабеля", "Прокладання кабелю"),
    "work.chase": (
        "work",
        "м",
        "Штроба узкая, 1–2 кабеля",
        "Штроба вузька, 1–2 кабелі",
    ),
    "work.chase_wide": (
        "work",
        "м",
        "Штроба широкая, 3–4 кабеля",
        "Штроба широка, 3–4 кабелі",
    ),
    "work.wall_hole": (
        "work",
        "шт",
        "Сверление прохода через стену",
        "Свердління проходу крізь стіну",
    ),
    "work.socket_box": (
        "work",
        "шт",
        "Сверление и установка подрозетника",
        "Свердління і встановлення підрозетника",
    ),
    "work.junction_box": (
        "work",
        "шт",
        "Монтаж коробки с опрессовкой",
        "Монтаж коробки з опресуванням",
    ),
    "work.board_install": ("work", "шт", "Установка щита", "Встановлення щита"),
    "work.board_module": ("work", "мод.", "Сборка щита", "Збирання щита"),
    "work.mechanism": ("work", "шт", "Установка механизмов", "Встановлення механізмів"),
    "work.marking": (
        "work",
        "шт",
        "Маркировка концов кабеля",
        "Маркування кінців кабелю",
    ),
    "work.insulation_test": (
        "work",
        "гр.",
        "Замер сопротивления изоляции с протоколом",
        "Вимірювання опору ізоляції з протоколом",
    ),
}
# Пояснения к позициям — значок «i» у наименования в просмотре, комментарий в Excel; без цены
# — ещё «цена уточняется»
HINTS: dict[str, dict[str, str]] = {
    "work.mechanism": {
        "ru": "Установка розеток, выключателей, интернет- и ТВ-выводов",
        "uk": "Встановлення розеток, вимикачів, інтернет- і ТВ-виводів",
    },
}
PRICE_TBD = {"ru": "Цена уточняется", "uk": "Ціна уточнюється"}
PARAM_KEYS = ("cable", "frame", "junction_box", "junction_box_ip", *(k.value for k in DeviceKind))
# Светильники и выводы: ключ → название в texts (как в условных обозначениях)
FIXTURES = {
    "light": "symbol_light",
    "downlight": "symbol_downlight",
    "pendant": "symbol_pendant",
    "sconce": "symbol_sconce",
    "led_output": "symbol_led_output",
    "outlet": "symbol_outlet",
    "led_strip": "legend_led_strip",
}
TEXT: dict[str, dict[str, str]] = {
    "material": {"ru": "Материалы", "uk": "Матеріали"},
    "work": {"ru": "Работы", "uk": "Роботи"},
    "head": {
        "ru": "| № | Наименование | Ед. | Кол. | Цена, {c} | Сумма, {c} |",
        "uk": "| № | Найменування | Од. | Кільк. | Ціна, {c} | Сума, {c} |",
    },
    "subtotal": {"ru": "Итого", "uk": "Разом"},
    "total": {"ru": "Всего", "uk": "Усього"},
    "no_price": {"ru": "нет цены", "uk": "немає ціни"},
    "unpriced": {
        "ru": "Без цены: {n} поз. — в итог не входят, нужна цена в {f}.",
        "uk": "Без ціни: {n} поз. — до підсумку не входять, потрібна ціна в {f}.",
    },
    "title": {"ru": "Смета электрики: {name}", "uk": "Кошторис електрики: {name}"},
    "sheet_title": {"ru": "Смета электрики", "uk": "Кошторис електрики"},
    "sheet_volumes": {"ru": "Ведомость объёмов электрики", "uk": "Відомість обсягів електрики"},
    "col_no": {"ru": "№", "uk": "№"},
    "col_name": {"ru": "Наименование", "uk": "Найменування"},
    "col_unit": {"ru": "Ед.", "uk": "Од."},
    "col_qty": {"ru": "Кол.", "uk": "Кільк."},
    "col_price": {"ru": "Цена, {c}", "uk": "Ціна, {c}"},
    "col_sum": {"ru": "Сумма, {c}", "uk": "Сума, {c}"},
    "section_total": {"ru": "Итого: {s}", "uk": "Разом: {s}"},
    "sheet_prices": {"ru": "Цены: {f}{date}", "uk": "Ціни: {f}{date}"},
    "note_prices": {"ru": "Цены: {f}{date}{note}", "uk": "Ціни: {f}{date}{note}"},
    "note_no_prices": {
        "ru": "Цен нет: {f} не найден — смета только с объёмами",
        "uk": "Цін немає: {f} не знайдено — кошторис лише з обсягами",
    },
    "note_cable": {
        "ru": "Кабель — с запасом {pct} %; прокладка — по фактической длине",
        "uk": "Кабель — із запасом {pct} %; прокладання — за фактичною довжиною",
    },
    "note_chases": {
        "ru": "Штробы: общая, если спуски ближе {d} мм, до {k} кабелей; до {n} кабелей — "
        "узкая, больше — широкая; силовые и слаботочка — в разных штробах",
        "uk": "Штроби: спільна, якщо спуски ближче {d} мм, до {k} кабелів; до {n} кабелів — "
        "вузька, більше — широка; силові і слабкострумові — в різних штробах",
    },
    "note_holes": {
        "ru": "Проходы через стены — над дверями и проёмами для трасс по потолку, до {k} "
        "кабелей в отверстии; трассы по полу идут под проёмом без сверления",
        "uk": "Проходи крізь стіни — над дверима і прорізами для трас по стелі, до {k} "
        "кабелів в отворі; траси по підлозі йдуть під прорізом без свердління",
    },
    "note_board_tail": {
        "ru": "Выпуск кабеля в щите — до пола ({h} мм), на остальных подключениях — {e} мм",
        "uk": "Виведення кабелю в щиті — до підлоги ({h} мм), на інших підключеннях — {e} мм",
    },
    "note_direct": {
        "ru": "Отдельные линии от щита без коробок: {groups}",
        "uk": "Окремі лінії від щита без коробок: {groups}",
    },
    "note_wet": {
        "ru": "Коробки во влажных помещениях — IP44: {n} шт",
        "uk": "Коробки у вологих приміщеннях — IP44: {n} шт",
    },
    "note_fixing": {
        "ru": "Крепления кабеля — шаг {step} мм",
        "uk": "Кріплення кабелю — крок {step} мм",
    },
    "note_crimp": {
        "ru": "Опрессовка: в каждой коробке L, N, PE и по соединению на клавишу выключателя",
        "uk": "Опресування: у кожній коробці L, N, PE і по з'єднанню на клавішу вимикача",
    },
    "note_excluded": {
        "ru": "Не входит: профиль и блоки питания LED, монтаж светильников, ввод питания",
        "uk": "Не входить: профіль і блоки живлення LED, монтаж світильників, введення живлення",
    },
    "note_panel": {
        "ru": "Щит — по составу листа «Щит»: аппараты по видам и номиналам",
        "uk": "Щит — за складом аркуша «Щит»: апарати за видами і номіналами",
    },
    "note_board_kit": {
        "ru": "board_kit в {f} не используется: состав щита — с листа «Щит», цены аппаратов — "
        "ключами {keys}",
        "uk": "board_kit у {f} не використовується: склад щита — з аркуша «Щит», ціни апаратів — "
        "ключами {keys}",
    },
    "note_height": {
        "ru": "Высота потолка 2700 мм — допущение (в room.yaml не указана)",
        "uk": "Висота стелі 2700 мм — припущення (у room.yaml не вказана)",
    },
    "note_questions": {
        "ru": "Открытые вопросы по электрике: {ids} — объёмы могут измениться",
        "uk": "Відкриті питання з електрики: {ids} — обсяги можуть змінитися",
    },
}


def text(key: str, language: Language, **values: object) -> str:
    return TEXT[key][language].format(**values)


class Price(StrictModel):
    price: float = Field(ge=0)
    name: str | None = Field(default=None, description="Своё наименование вместо стандартного")


class BoardItem(StrictModel):
    """Устарело: состав щита считает лист «Щит», цены аппаратов — ключами по видам
    (breaker, diff, rcd, input…). Поле читается, чтобы старые прайсы открывались."""

    name: str = Field(min_length=1)
    price: float = Field(ge=0)
    qty: int = Field(default=1, ge=1)
    modules: int = Field(default=0, ge=0, description="Ширина на DIN-рейке, модулей")


class PriceList(StrictModel):
    """prices.yaml: цены по ключам позиций (LINES), за единицу, в валюте currency."""

    currency: str = "грн"
    date: str = Field(default="", description="Когда собраны цены")
    note: str = ""
    items: dict[str, float | Price] = Field(default_factory=dict)
    board_kit: list[BoardItem] = Field(
        default_factory=list, description="Устарело — не используется"
    )

    def price(self, key: str) -> Price | None:
        item = self.items.get(key)
        return Price(price=item) if isinstance(item, int | float) else item


def split_key(key: str) -> tuple[str, str]:
    """«cable.3x1.5» → («cable», «3x1.5»); ключ без параметра → (ключ, «»)."""
    base, _, param = key.partition(".")
    return (base, param) if base in PARAM_KEYS and param else (key, "")


def unknown_keys(prices: PriceList) -> list[str]:
    """Ключи прайса, которых нет среди позиций сметы (опечатка или чужой ключ)."""
    return [k for k in prices.items if split_key(k)[0] not in LINES]


@dataclass(frozen=True)
class Line:
    section: Section
    key: str
    name: str
    unit: str
    qty: float
    price: float | None  # None — в прайсе нет
    hint: str | None = None  # пояснение к позиции (HINTS)

    @property
    def total(self) -> float | None:
        return None if self.price is None else self.qty * self.price


@dataclass(frozen=True)
class Estimate:
    lines: tuple[Line, ...]
    currency: str

    def subtotal(self, section: Section) -> float:
        return sum(ln.total or 0.0 for ln in self.lines if ln.section == section)

    @property
    def total(self) -> float:
        return self.subtotal("material") + self.subtotal("work")

    @property
    def unpriced(self) -> int:
        return sum(1 for ln in self.lines if ln.price is None)


def quantities(wiring: Wiring, elec: Electrical) -> dict[str, float]:
    """Объёмы по ключам позиций (без цен)."""
    q: dict[str, float] = {}
    for mark, m in wiring.by_mark().items():
        q[f"cable.{mark}"] = round(m * (1 + RESERVE), 1)
    mounted = [p for p in elec.points if p.type in MOUNTED]
    posts = sum(p.count for p in mounted)
    q["socket_box"] = posts
    for b in sorted(wiring.boxes, key=lambda b: BOX_ORDER.index(b.size)):
        key = f"{'junction_box_ip' if b.wet else 'junction_box'}.{b.size}"
        q[key] = q.get(key, 0) + 1
    by_id = {p.id: p for p in elec.points}
    keys = sum(
        by_id[r.target].keys
        for r in wiring.runs
        if r.source != "board" and r.target in by_id and by_id[r.target].type in SWITCHES
    )
    q["crimp"] = 3 * len(wiring.boxes) + keys
    q["cable_fixing"] = math.ceil(sum(r.plan_mm for r in wiring.runs) / FIXING_STEP)
    if elec.board is not None:
        q["board_enclosure"] = 1
    q["breaker"] = len({r.group for r in wiring.runs})
    pieces: Counter[str] = Counter()
    for p in mounted:
        if p.type is PointType.SWITCH:
            pieces["switch2" if p.keys == 2 else "switch1"] += p.count
        else:
            pieces[p.type.value] += p.count
    for key in ("socket", "switch1", "switch2", "pass_switch", "data", "tv"):
        if pieces[key]:
            q[key] = pieces[key]
    for n, k in sorted(Counter(p.count for p in mounted).items()):
        q[f"frame.{n}"] = k
    fixtures: Counter[str] = Counter()
    for p in elec.points:
        if p.type.value in FIXTURES:
            fixtures[p.type.value] += p.count
    for key in FIXTURES:
        if fixtures[key]:
            q[key] = fixtures[key]
    led_m = sum(s.end - s.start for s in elec.led_strips) / 1000
    if led_m:
        q["led_strip"] = round(led_m, 2)
    q["work.cable"] = round(sum(r.length_mm for r in wiring.runs) / 1000, 1)
    q["work.chase"] = round(wiring.chase_length(wide=False) / 1000, 1)
    q["work.chase_wide"] = round(wiring.chase_length(wide=True) / 1000, 1)
    q["work.wall_hole"] = len(wiring.holes)
    q["work.socket_box"] = posts
    q["work.junction_box"] = len(wiring.boxes)
    if elec.board is not None:
        q["work.board_install"] = 1
    q["work.mechanism"] = sum(pieces.values())
    q["work.marking"] = 2 * len(wiring.runs)
    q["work.insulation_test"] = len({r.group for r in wiring.runs if r.circuit == "power"})
    return {k: v for k, v in q.items() if v}


def build_estimate(
    wiring: Wiring,
    elec: Electrical,
    prices: PriceList,
    language: Language = "ru",
    panel: BoardPanel | None = None,
) -> Estimate:
    """Смета: позиция на каждый объём. Щит — по составу листа «Щит» (panel): строка на вид и
    номинал аппарата, сборка — по модулям состава; щита нет — автомат на группу."""
    q = quantities(wiring, elec)
    if panel is not None:
        q.pop("breaker", None)
    if "work.board_install" in q:  # сборка — сразу за установкой щита
        modules = panel.modules if panel is not None else q.get("breaker", 0.0)
        items = list(q.items())
        at = next(k for k, (key, _) in enumerate(items) if key == "work.board_install") + 1
        items.insert(at, ("work.board_module", modules))
        q = dict(items)

    def price(key: str) -> Price | None:
        base, param = split_key(key)
        return prices.price(key) or (prices.price(base) if param else None)

    def line(key: str, qty: float, name: str | None = None) -> Line:
        base, param = split_key(key)
        section, unit, ru, uk = LINES[base]
        p = price(key)
        exact = prices.price(key)
        if exact and exact.name:
            name = exact.name
        elif p and p.name and name is None:
            name = p.name
        elif name is None and base in FIXTURES:
            name = tr(FIXTURES[base], language)
        elif name is None:
            name = (ru if language == "ru" else uk).format(
                mark=BOX_TITLES.get(param, param), n=param
            )
        hint = HINTS.get(base, {}).get(language)
        if hint and p is None:
            hint = f"{hint}. {PRICE_TBD[language]}"
        return Line(section, key, name, unit, qty, p.price if p else None, hint)

    lines = [line(k, v) for k, v in q.items()]
    if panel is not None:
        # Аппараты щита — по видам и номиналам, в порядке щита; рядом с корпусом
        counts: dict[tuple[str, str], tuple[str, int]] = {}
        for d in panel.devices:
            key = f"{d.kind.value}.{d.label}"
            name, n = counts.get((key, d.label), (device_name(d, language), 0))
            counts[(key, d.label)] = (name, n + 1)
        devices = [line(key, n, name) for (key, _), (name, n) in counts.items()]
        at = next((k + 1 for k, ln in enumerate(lines) if ln.key == "board_enclosure"), len(lines))
        lines[at:at] = devices
    lines.sort(key=lambda ln: 0 if ln.section == "material" else 1)
    return Estimate(tuple(lines), prices.currency)


def format_number(x: float, digits: int = 2) -> str:
    text = f"{x:,.{digits}f}".replace(",", " ").replace(".", ",")
    return text.removesuffix(",00") if digits else text


def write_markdown(
    est: Estimate, path: Path, title: str, notes: list[str], language: Language
) -> Path:
    def t(key: str, **values: object) -> str:
        return text(key, language, **values)

    out = [f"# {title}", ""]
    for section in ("material", "work"):
        rows = [ln for ln in est.lines if ln.section == section]
        if not rows:
            continue
        out += [f"## {t(section)}", "", t("head", c=est.currency), "|---|---|---|---|---:|---:|"]
        for k, ln in enumerate(rows, 1):
            price = format_number(ln.price) if ln.price is not None else t("no_price")
            total = format_number(ln.total) if ln.total is not None else "—"
            qty = format_number(ln.qty, 1).removesuffix(",0")
            out.append(f"| {k} | {ln.name} | {ln.unit} | {qty} | {price} | {total} |")
        out += [f"| | **{t('subtotal')}** | | | | **{format_number(est.subtotal(section))}** |", ""]
    out += [f"**{t('total')}: {format_number(est.total)} {est.currency}**", ""]
    if est.unpriced:
        out += [t("unpriced", n=est.unpriced, f=PRICES_FILE), ""]
    out += [f"- {n}" for n in notes]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(out) + "\n", encoding="utf-8")
    return path


def write_csv(est: Estimate, path: Path) -> Path:
    """CSV для Excel: UTF-8 с BOM, разделитель «;», десятичная запятая."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["section", "key", "name", "unit", "qty", "price", "total"])
        for ln in est.lines:
            w.writerow(
                [
                    ln.section,
                    ln.key,
                    ln.name,
                    ln.unit,
                    f"{ln.qty:g}".replace(".", ","),
                    "" if ln.price is None else f"{ln.price:g}".replace(".", ","),
                    "" if ln.total is None else f"{ln.total:.2f}".replace(".", ","),
                ]
            )
    return path


def write_xlsx(est: Estimate, path: Path, title: str, notes: list[str], language: Language) -> Path:
    """Смета в Excel: разделы, сумма строки — формулой «кол. × цена», итоги разделов и
    «Всего» — формулами, цены и количества подсвечены (их правят). Позиция без цены —
    пустая цена с примечанием; впишут цену — войдёт в итоги."""

    def t(key: str, **values: object) -> str:
        return text(key, language, **values)

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = t("sheet_title")[:31]
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    heads = [
        t("col_no"),
        t("col_name"),
        t("col_unit"),
        t("col_qty"),
        t("col_price", c=est.currency),
        t("col_sum", c=est.currency),
        "key",
    ]
    ws.append([])
    ws.append(heads)
    head_row = ws.max_row
    for cell in ws[head_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    edit = PatternFill("solid", start_color=XLSX_EDIT_FILL, end_color=XLSX_EDIT_FILL)
    grey = Font(color="999999", size=8)
    subtotals: list[str] = []
    for section in ("material", "work"):
        rows = [ln for ln in est.lines if ln.section == section]
        if not rows:
            continue
        ws.append(["", t(section)])
        ws.cell(ws.max_row, 2).font = Font(bold=True)
        first = ws.max_row + 1
        for k, ln in enumerate(rows, 1):
            r = ws.max_row + 1
            ws.append([k, ln.name, ln.unit, ln.qty, ln.price, f'=IF(E{r}="","",D{r}*E{r})', ln.key])
            ws.cell(r, 4).number_format = XLSX_QTY
            ws.cell(r, 5).number_format = XLSX_MONEY
            ws.cell(r, 6).number_format = XLSX_MONEY
            ws.cell(r, 4).fill = edit
            ws.cell(r, 5).fill = edit
            ws.cell(r, 7).font = grey
            if ln.hint:
                ws.cell(r, 2).comment = Comment(ln.hint, "Shtroba")
            if ln.price is None:
                ws.cell(r, 5).comment = Comment(t("no_price"), "Shtroba")
        last = ws.max_row
        ws.append(["", t("section_total", s=t(section)), "", "", "", f"=SUM(F{first}:F{last})"])
        total_row = ws.max_row
        ws.cell(total_row, 2).font = Font(bold=True)
        ws.cell(total_row, 6).font = Font(bold=True)
        ws.cell(total_row, 6).number_format = XLSX_MONEY
        subtotals.append(f"F{total_row}")
    ws.append([])
    ws.append(
        [
            "",
            f"{t('total')}, {est.currency}",
            "",
            "",
            "",
            "=" + "+".join(subtotals) if subtotals else 0,
        ]
    )
    grand = ws.max_row
    for col in (2, 6):
        ws.cell(grand, col).font = Font(bold=True, size=12)
    ws.cell(grand, 6).number_format = XLSX_MONEY
    ws.append([])
    if est.unpriced:
        ws.append(["", t("unpriced", n=est.unpriced, f=PRICES_FILE)])
    for note in notes:
        ws.append(["", note])
    for letter, width in zip("ABCDEFG", (5, 48, 7, 10, 12, 14, 18), strict=True):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = f"A{head_row + 1}"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
