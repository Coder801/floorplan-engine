"""Смета электрики: объёмы из модели, цены из прайса, позиции без цены."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
import yaml

from floorplan_engine.estimate import PriceList, build_estimate, quantities, unknown_keys
from floorplan_engine.pipeline import estimate_project
from floorplan_engine.schema import Electrical, Plan
from floorplan_engine.validate import validate_plan
from floorplan_engine.wiring import RESERVE, box_wiring
from tests.conftest import make_plan

RECT = [{"length": 3000}, {"length": 4000}, {"length": 3000}, {"length": 4000}]
DOOR = {"wall": "w2", "type": "door", "offset": 1000, "width": 900}


def _pt(
    pid: str, kind: str, offset: float, height: float, group: str, **extra: Any
) -> dict[str, Any]:
    return {
        "id": pid,
        "room": "r",
        "type": kind,
        "wall": "w3",
        "offset": offset,
        "height": height,
        "group": group,
        **extra,
    }


ELEC: dict[str, Any] = {
    "points": [
        {"id": "l1", "room": "r", "type": "downlight", "ceiling": [1000, 1500], "group": "QF1"},
        _pt("k1", "switch", 2500, 900, "QF1", keys=2),
        _pt("s1", "socket", 500, 300, "QF2", count=2),
        _pt("s2", "socket", 1500, 300, "QF2"),
    ],
    "board": {"room": "r", "wall": "w1", "offset": 500, "height": 1700},
}


def _wiring() -> tuple[Any, Electrical]:
    plan = make_plan(RECT, [DOOR])
    _, contours = validate_plan(plan)
    elec = Electrical.model_validate(ELEC)
    wiring = box_wiring(plan, contours, elec)
    assert wiring is not None
    return wiring, elec


def test_quantities_from_model() -> None:
    wiring, elec = _wiring()
    q = quantities(wiring, elec)
    for mark, m in wiring.by_mark().items():
        assert q[f"cable.{mark}"] == pytest.approx(round(m * (1 + RESERVE), 1))
    assert q["socket_box"] == 4  # 2 + 1 розетки + выключатель
    assert q["frame.1"] == 2 and q["frame.2"] == 1
    assert q["switch2"] == 1 and q["socket"] == 3
    boxes = len(wiring.boxes)
    assert sum(v for k, v in q.items() if k.startswith("junction_box.")) == boxes
    assert all(f"junction_box.{b.size}" in q for b in wiring.boxes)  # по типоразмерам
    assert q["crimp"] == 3 * boxes + 2  # двухклавишный выключатель от коробки
    assert q["breaker"] == 2
    narrow, wide = wiring.chase_length(wide=False), wiring.chase_length(wide=True)
    assert narrow + wide == pytest.approx(wiring.chase_mm)
    assert q["work.chase"] == pytest.approx(round(narrow / 1000, 1))
    assert q.get("work.chase_wide", 0) == pytest.approx(round(wide / 1000, 1))
    assert q["work.marking"] == 2 * len(wiring.runs)  # бирка на каждом конце
    assert q["work.insulation_test"] == 2  # по силовым группам
    assert q["work.cable"] == pytest.approx(round(sum(r.length_mm for r in wiring.runs) / 1000, 1))


def test_prices_and_unpriced_lines() -> None:
    wiring, elec = _wiring()
    prices = PriceList.model_validate(
        {
            "items": {
                "cable.3x2.5": 70,
                "work.board_module": 100,
                "socket_box": {"price": 20, "name": "Підрозетник Kopos"},
            },
            "board_kit": [{"name": "Реле", "price": 1500, "modules": 2}],
        }
    )
    est = build_estimate(wiring, elec, prices, "uk")
    by_key = {ln.key: ln for ln in est.lines}
    assert by_key["socket_box"].name == "Підрозетник Kopos" and by_key["socket_box"].total == 80
    box = next(ln for k, ln in by_key.items() if k.startswith("junction_box."))
    assert box.price is None and est.unpriced > 0
    module = by_key["work.board_module"]
    assert module.qty == 2 + 2 and module.total == 400  # 2 автомата + реле на 2 модуля
    assert est.total == pytest.approx(sum(ln.total or 0 for ln in est.lines))
    sections = [ln.section for ln in est.lines]
    assert sections == sorted(sections, key=lambda s: s != "material")  # сначала материалы
    assert unknown_keys(PriceList.model_validate({"items": {"cabel.3x1.5": 1, "frame.2": 5}})) == [
        "cabel.3x1.5"
    ]


def test_estimate_project_writes_tables(tmp_path: Path) -> None:
    project = tmp_path / "kv1"
    project.mkdir()
    room = {
        "name": "Кв",
        "default_wall_thickness": 120,
        "source": "text",
        "rooms": [
            {
                "id": "r",
                "name": "К",
                "walls": [{"id": f"w{i + 1}", **w} for i, w in enumerate(RECT)],
                "openings": [{"id": "r_door_1", **DOOR}],
            }
        ],
    }
    (project / "room.yaml").write_text(yaml.safe_dump(room, allow_unicode=True), "utf-8")
    (project / "electrical.yaml").write_text(yaml.safe_dump(ELEC, allow_unicode=True), "utf-8")

    no_prices = estimate_project(project)
    assert no_prices.estimate is not None and no_prices.estimate.total == 0
    assert no_prices.estimate.unpriced == len(no_prices.estimate.lines)

    (project / "prices.yaml").write_text("items:\n  socket_box: 20\n  work.chase: 200\n", "utf-8")
    result = estimate_project(project)
    assert result.estimate is not None and result.estimate.total > 0
    md, csv_path, xlsx = result.files
    assert xlsx.suffix == ".xlsx" and xlsx.stat().st_size > 0
    assert "Смета электрики: Кв" in md.read_text("utf-8")
    raw = csv_path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM — Excel узнаёт UTF-8
    assert "socket_box;Подрозетник;шт;4;20;80,00" in raw.decode("utf-8-sig")

    (project / "prices.yaml").write_text("items:\n  sockt_box: 20\n", "utf-8")
    bad = estimate_project(project)
    assert bad.estimate is None
    assert any(i.file == "prices.yaml" and i.path == "items.sockt_box" for i in bad.issues)


def test_wet_boxes_and_wall_holes_in_estimate() -> None:
    data = make_plan(RECT, [DOOR]).model_dump()
    data["rooms"][0]["wet"] = True
    plan = Plan.model_validate(data)
    _, contours = validate_plan(plan)
    elec = Electrical.model_validate(ELEC)
    wiring = box_wiring(plan, contours, elec)
    assert wiring is not None
    q = quantities(wiring, elec)
    ip = sum(v for k, v in q.items() if k.startswith("junction_box_ip."))
    assert ip == len(wiring.boxes) and not any(k.startswith("junction_box.") for k in q)
    assert q.get("work.wall_hole", 0) == len(wiring.holes)


def test_estimate_table_with_and_without_prices() -> None:
    from floorplan_engine.estimate import estimate_table

    wiring, elec = _wiring()
    prices = PriceList.model_validate({"items": {"socket_box": 20}})
    priced = build_estimate(wiring, elec, prices, "uk")
    table = estimate_table(priced, "uk", priced=True)
    assert table.title == "Кошторис електрики" and len(table.columns) == 6
    names = [r[1] for r in table.rows]
    assert names[0] == "Матеріали" and "Роботи" in names
    assert "Разом: матеріали" in names and "Разом: роботи" in names
    assert table.total is not None and table.total[1] == "Усього"
    socket_box = next(r for r in table.rows if r[1] == "Підрозетник")
    assert socket_box[4] == "20" and socket_box[5] == "80"  # 4 поста × 20
    unpriced = next(r for r in table.rows if r[1] == "Розетка із заземленням")
    assert unpriced[4] == unpriced[5] == "—"

    volumes = estimate_table(build_estimate(wiring, elec, PriceList(), "uk"), "uk", priced=False)
    assert volumes.title == "Відомість обсягів електрики"
    assert len(volumes.columns) == 4 and volumes.total is None  # без цен и итогов
    assert not any(str(r[1]).startswith("Разом") for r in volumes.rows)


def test_estimate_on_wiring_sheet(tmp_path: Path) -> None:
    """Смета стоит на CAD-листе проводки — тем же текстом, что в таблице."""
    from ezdxf.filemanagement import readfile

    from floorplan_engine.preview import sheet_layouts
    from floorplan_engine.sheets import build_project_document

    plan = make_plan(RECT, [DOOR])
    _, contours = validate_plan(plan)
    elec = Electrical.model_validate(ELEC)
    prices = PriceList.model_validate({"items": {"socket_box": 20}})
    path = tmp_path / "project.dxf"
    build_project_document(plan, contours, elec, prices=prices).saveas(path, encoding="utf-8")
    texts = {
        lay.name: {e.dxf.text for e in lay if e.dxftype() == "TEXT"}
        for lay in sheet_layouts(readfile(path))
    }
    wiring_tab = next(name for name in texts if name.endswith("Проводка"))
    assert "Смета электрики" in texts[wiring_tab]  # язык плана по умолчанию — ru
    assert not any("Смета электрики" in t for name, t in texts.items() if name != wiring_tab)


def test_box_price_falls_back_to_size_free_key() -> None:
    wiring, elec = _wiring()
    prices = PriceList.model_validate({"items": {"junction_box": 60}})
    est = build_estimate(wiring, elec, prices, "uk")
    boxes = [ln for ln in est.lines if ln.key.startswith("junction_box.")]
    assert boxes and all(ln.price == 60 for ln in boxes)
    assert all("×" in ln.name or "Ø" in ln.name for ln in boxes)  # размер в наименовании


def _eval_xlsx(path: Path) -> float:
    """Пересчитать «Всего» книги сметы по её формулам (Excel под рукой нет)."""
    import re

    from openpyxl import load_workbook

    ws = load_workbook(path).active
    assert ws is not None
    values: dict[str, float] = {}

    def cell(ref: str) -> float:
        if ref in values:
            return values[ref]
        v = ws[ref].value
        if isinstance(v, str) and v.startswith("=IF("):
            r = re.findall(r"\d+", v)[0]
            price = ws[f"E{r}"].value
            out = 0.0 if price is None else float(ws[f"D{r}"].value) * float(price)
        elif isinstance(v, str) and v.startswith("=SUM("):
            a, b = (int(x) for x in re.findall(r"F(\d+)", v))
            out = sum(cell(f"F{k}") for k in range(a, b + 1))
        elif isinstance(v, str) and v.startswith("="):
            out = sum(cell(ref) for ref in re.findall(r"F\d+", v))
        else:
            out = float(v or 0)
        values[ref] = out
        return out

    grand = next(
        r for r in range(1, ws.max_row + 1) if str(ws[f"B{r}"].value or "").startswith("Усього")
    )
    return cell(f"F{grand}")


def test_xlsx_formulas_give_estimate_total(tmp_path: Path) -> None:
    wiring, elec = _wiring()
    prices = PriceList.model_validate({"items": {"cable.3x2.5": 70, "work.cable": 30}})
    est = build_estimate(wiring, elec, prices, "uk")
    from floorplan_engine.estimate import write_xlsx

    path = write_xlsx(est, tmp_path / "e.xlsx", "Кошторис", ["примітка"], "uk")
    assert _eval_xlsx(path) == pytest.approx(est.total)
    assert est.unpriced > 0  # позиции без цены — в книге, но в итог не входят


def test_studio_prices_used_when_project_has_none(tmp_path: Path) -> None:
    from floorplan_engine.pipeline import find_prices

    project = tmp_path / "kv"
    project.mkdir()
    assert find_prices(project) is None
    (tmp_path / "prices.yaml").write_text("items: {}\n", "utf-8")
    assert find_prices(project) == tmp_path / "prices.yaml"
    (project / "prices.yaml").write_text("items: {}\n", "utf-8")
    assert find_prices(project) == project / "prices.yaml"  # свой прайс важнее студии
